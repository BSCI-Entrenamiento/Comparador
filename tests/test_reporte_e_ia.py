import json
from types import SimpleNamespace

import anthropic
import openpyxl

import config
from comparador import emparejador as em
from comparador import reporte
from comparador.modelos import Producto
from comparador.revisor_ia import Decision, RevisorIA, aplicar


def _resultados(catalogos):
    im, comp = catalogos
    return {p.url: {t: em.emparejar(p, comp[t], t) for t in config.TIENDAS_COMPETENCIA} for p in im.values()}


def test_excel_y_historico(catalogos, tmp_path):
    im, comp = catalogos
    resultados = _resultados(catalogos)
    ruta = reporte.generar_excel(tmp_path / "r.xlsx", "laptops", list(im.values()), resultados, comp,
                                 {"Click": "prueba de tienda caída"}, "IA: no usada")
    wb = openpyxl.load_workbook(ruta)
    assert wb.sheetnames == ["Comparación", "Resumen", "Detalle", "Catálogos"]
    ws = wb["Comparación"]
    assert ws.max_row == 2 + len(im)
    assert "NO DISPONIBLE" in ws.cell(row=1, column=18).value  # bloque de Click (3a tienda)
    reporte.agregar_historico(tmp_path / "h.csv", list(im.values()), resultados)
    lineas = (tmp_path / "h.csv").read_text(encoding="utf-8-sig").splitlines()
    assert len(lineas) == 1 + len(im) * len(config.TIENDAS_COMPETENCIA)


def test_precio_comparable_intelaf(monkeypatch):
    promo = Producto("Intelaf", "x", 8011.0, "u", precio_alterno=6499.0, nota_precio="Promo ASUS")
    efectivo = Producto("Intelaf", "x", 9941.0, "u", precio_alterno=9245.0, nota_precio="Beneficio Efectivo")
    monkeypatch.setattr(config, "INTELAF_PRECIO", "promo")
    assert reporte.precio_comparable(promo) == 6499.0
    assert reporte.precio_comparable(efectivo) == 9941.0
    monkeypatch.setattr(config, "INTELAF_PRECIO", "oferta")
    assert reporte.precio_comparable(efectivo) == 9245.0
    monkeypatch.setattr(config, "INTELAF_PRECIO", "normal")
    assert reporte.precio_comparable(promo) == 8011.0


class _ClienteFalso:
    """Imita client.beta.messages.create y cuenta las llamadas."""

    def __init__(self, decision: dict):
        self.llamadas = 0
        self.decision = decision
        self.beta = SimpleNamespace(messages=SimpleNamespace(create=self._create))

    def _create(self, **kwargs):
        self.llamadas += 1
        assert kwargs["output_config"]["format"]["type"] == "json_schema"
        return SimpleNamespace(
            stop_reason="end_turn", stop_details=None,
            usage=SimpleNamespace(input_tokens=1200, output_tokens=80),
            content=[SimpleNamespace(type="text", text=json.dumps(self.decision))],
        )


def _revisor(tmp_path, decision):
    r = RevisorIA(tmp_path / "cache_ia.json")
    r._anthropic, r.cliente, r.activo = anthropic, _ClienteFalso(decision), True
    return r


def test_ia_solo_revisa_dudosos_y_usa_cache(catalogos, tmp_path):
    im, comp = catalogos
    p = im["83S0001KGJ"]  # Click: "Similar (mismo modelo, otra config.)"
    c = em.emparejar(p, comp["Click"], "Click")
    decision = {"candidato": 1, "nivel": "equivalente", "diferencias": "Click trae 512 GB", "confianza": 0.9}
    r = _revisor(tmp_path, decision)
    nuevo = r.revisar(p, "Click", c)
    assert nuevo.nivel == em.SIMILAR_IA and "512 GB" in nuevo.nota_ia
    assert r.cliente.llamadas == 1

    r2 = _revisor(tmp_path, decision)  # nueva corrida: sale de caché
    r2.revisar(p, "Click", em.emparejar(p, comp["Click"], "Click"))
    assert r2.cliente.llamadas == 0 and r2.desde_cache == 1

    exacto = em.emparejar(im["16-BY0015DX"], comp["Click"], "Click")
    assert r.revisar(im["16-BY0015DX"], "Click", exacto) is exacto  # los exactos no se consultan
    assert r.cliente.llamadas == 1


def test_ia_no_puede_dar_exacto_de_otra_marca(catalogos):
    im, comp = catalogos
    p = im["LDC15250-5434BLK-PUS"]
    c = em.emparejar(p, comp["MacroSistemas"], "MacroSistemas")  # candidato 1: ASUS
    d = Decision(candidato=1, nivel="exacto", diferencias="misma configuración", confianza=0.9)
    assert aplicar(c, d, c.candidatos, p).nivel == em.SIMILAR_IA


def test_ia_no_puede_saltarse_el_criterio(catalogos):
    im, comp = catalogos
    p = im["83NM8"]  # Dell 15 DC15250 con 8GB
    c = em.emparejar(p, comp["Intelaf"], "Intelaf")
    i = next(i for i, (_, _, x, _, _) in enumerate(c.candidatos, 1) if x.specs["ram"] == 16)
    d = Decision(candidato=i, nivel="equivalente", diferencias="trae 16GB", confianza=0.9)
    nuevo = aplicar(c, d, c.candidatos, p)
    assert nuevo.producto is None and nuevo.nivel == em.SIN
    assert "otra RAM" in nuevo.nota_ia


def test_ia_otra_marca_requiere_misma_configuracion(catalogos):
    im, comp = catalogos
    p = im["83S0001KGJ"]  # Lenovo LOQ 1TB
    msi = next(x for x in comp["Intelaf"] if "Thin A15" in x.nombre)  # MSI con 512GB
    c = em.emparejar(p, comp["Intelaf"], "Intelaf")
    candidatos = [(em.SIN, 31.0, msi, "parcial", "")]
    d = Decision(candidato=1, nivel="equivalente", diferencias="otra marca, 512GB", confianza=0.7)
    nuevo = aplicar(c, d, candidatos, p)
    assert nuevo.producto is None and "otra marca con otra configuración" in nuevo.nota_ia


def test_ia_exacto_tactil_vs_no_tactil_baja_a_similar(catalogos):
    im, comp = catalogos
    p = im["B9TR0LA-ABM"]  # HP 15-fc027la, no táctil
    tactil = next(x for x in comp["Click"] if "15-FC0027CA" in x.nombre)  # versión táctil
    assert tactil.specs["tactil"] and not p.specs["tactil"]
    c = em.emparejar(p, comp["Click"], "Click")
    d = Decision(candidato=1, nivel="exacto", diferencias="misma config.", confianza=0.95)
    nuevo = aplicar(c, d, [(em.SIMILAR_SPECS, 75.0, tactil, "specs", "")], p)
    assert nuevo.nivel == em.SIMILAR_IA and "táctil" in nuevo.nota_ia


def test_cambio_de_criterio_invalida_la_cache(catalogos, tmp_path, monkeypatch):
    from comparador import revisor_ia

    im, comp = catalogos
    p = im["83S0001KGJ"]
    decision = {"candidato": 1, "nivel": "equivalente", "diferencias": "x", "confianza": 0.9}
    _revisor(tmp_path, decision).revisar(p, "Click", em.emparejar(p, comp["Click"], "Click"))
    monkeypatch.setattr(revisor_ia, "SISTEMA", revisor_ia.SISTEMA + " (criterio nuevo)")
    r = _revisor(tmp_path, decision)
    r.revisar(p, "Click", em.emparejar(p, comp["Click"], "Click"))
    assert r.cliente.llamadas == 1 and r.desde_cache == 0


def test_ia_descarta(catalogos):
    im, comp = catalogos
    c = em.emparejar(im["83S0001KGJ"], comp["Click"], "Click")
    d = Decision(candidato=None, nivel="distinto", diferencias="otro segmento", confianza=0.8)
    nuevo = aplicar(c, d, c.candidatos)
    assert nuevo.producto is None and nuevo.nivel == em.SIN
