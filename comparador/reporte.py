"""Generación del Excel de comparación y del histórico de precios."""
import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from comparador import emparejador as em
from comparador.modelos import Coincidencia, Producto

FMT_Q = '"Q"#,##0.00'
FMT_PCT = "0.0%"
AZUL = PatternFill("solid", fgColor="1F3864")
GRIS_TIENDA = PatternFill("solid", fgColor="D9E1F2")
VERDE = PatternFill("solid", fgColor="C6EFCE")
ROJO = PatternFill("solid", fgColor="FFC7CE")
AMARILLO = PatternFill("solid", fgColor="FFF2CC")
BLANCO_NEGRITA = Font(bold=True, color="FFFFFF")
GRIS_ITALICA = Font(italic=True, color="7F7F7F")
LINK = Font(color="0563C1", underline="single")


def _usa_oferta(p: Producto) -> bool:
    if p.tienda != "Intelaf" or not p.precio_alterno or config.INTELAF_PRECIO == "normal":
        return False
    if config.INTELAF_PRECIO == "oferta":
        return True
    nota = (p.nota_precio or "").upper()
    return not any(f in nota for f in config.FORMAS_DE_PAGO)


def precio_comparable(p: Producto) -> float | None:
    """Precio con el que se compara (ver config.INTELAF_PRECIO)."""
    return p.precio_alterno if _usa_oferta(p) else p.precio


def _detalle_precio(p: Producto) -> str:
    partes = []
    if p.tienda == "Intelaf" and p.precio_alterno:
        if _usa_oferta(p):
            partes.append(f"{p.nota_precio or 'Oferta'} (precio normal Q{p.precio:,.0f})")
        else:
            partes.append(f"{p.nota_precio or 'Oferta'}: Q{p.precio_alterno:,.0f}")
    if p.precio_lista:
        partes.append(f"Antes Q{p.precio_lista:,.0f}")
    if p.disponible is False:
        partes.append("Agotado")
    return "; ".join(partes)


def _encabezado(ws, fila: int, valores: list, relleno=AZUL, fuente=BLANCO_NEGRITA):
    for col, v in enumerate(valores, 1):
        celda = ws.cell(row=fila, column=col, value=v)
        celda.fill, celda.font = relleno, fuente
        celda.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _anchos(ws, anchos: dict):
    for col, ancho in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = ancho


def _link(celda, url: str):
    celda.hyperlink = url
    celda.font = LINK


def _specs_txt(s: dict) -> str:
    def gb(v):
        return "?" if v is None else (f"{v // 1024}TB" if v >= 1024 else f"{v}GB")
    return (f"{s.get('marca') or '?'} | {s.get('cpu') or '?'} | RAM {gb(s.get('ram'))} | "
            f"{gb(s.get('almacenamiento'))} | {s.get('pantalla') or '?'}\" | {s.get('gpu') or '-'}")


def _categoria(p: Producto) -> str:
    return p.categoria_nombre or p.categoria or ""


def _posicion(precio_im: float, mejores: list[float]) -> str:
    if not mejores:
        return "Sin comparación"
    minimo = min(mejores)
    if abs(precio_im - minimo) < 0.5:
        return "Igual al más barato"
    return "Más barato" if precio_im < minimo else "Más caro"


def generar_excel(ruta: Path, alcance: str, imeqmo: list[Producto],
                  resultados: dict[str, dict[str, Coincidencia]], catalogos: dict[str, list[Producto]],
                  errores: dict[str, str], resumen_ia: str) -> Path:
    tiendas = config.TIENDAS_COMPETENCIA
    imeqmo = sorted(imeqmo, key=lambda p: (_categoria(p), p.nombre))
    wb = Workbook()

    # ---------------------------------------------------------------- Comparación
    ws = wb.active
    ws.title = "Comparación"
    fijas = ["Categoría", "Producto Imeqmo", "No. parte", "UPC", "Precio Imeqmo"]
    por_tienda = ["Precio", "Imeqmo − comp. (Q)", "Imeqmo − comp. (%)", "Coincidencia", "Producto competidor", "Detalle"]
    finales = ["Mín. competencia (exactos)", "Tienda más barata", "Posición Imeqmo"]
    col = len(fijas) + 1
    for t in tiendas:
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(por_tienda) - 1)
        c = ws.cell(row=1, column=col, value=t + (f" — NO DISPONIBLE: {errores[t]}" if t in errores else ""))
        c.fill, c.font, c.alignment = GRIS_TIENDA, Font(bold=True), Alignment(horizontal="center")
        col += len(por_tienda)
    _encabezado(ws, 2, fijas + por_tienda * len(tiendas) + finales)

    for fila, p in enumerate(imeqmo, 3):
        ws.cell(row=fila, column=1, value=_categoria(p))
        _link(ws.cell(row=fila, column=2, value=p.nombre), p.url)
        ws.cell(row=fila, column=3, value=p.no_parte)
        ws.cell(row=fila, column=4, value=p.upc)
        ws.cell(row=fila, column=5, value=p.precio).number_format = FMT_Q
        col = len(fijas) + 1
        exactos = []
        for t in tiendas:
            c = resultados.get(p.url, {}).get(t)
            if c and c.producto:
                comp = c.producto
                precio = precio_comparable(comp)
                es_exacto = c.nivel in em.EXACTOS
                celdas = [
                    ws.cell(row=fila, column=col, value=precio),
                    ws.cell(row=fila, column=col + 1, value=(p.precio - precio) if precio else None),
                    ws.cell(row=fila, column=col + 2, value=((p.precio - precio) / precio) if precio else None),
                    ws.cell(row=fila, column=col + 3, value=c.nivel),
                    ws.cell(row=fila, column=col + 4, value=comp.nombre),
                    ws.cell(row=fila, column=col + 5, value="; ".join(
                        x for x in (c.diferencias, c.nota_ia and f"IA: {c.nota_ia}", _detalle_precio(comp)) if x)),
                ]
                celdas[0].number_format = celdas[1].number_format = FMT_Q
                celdas[2].number_format = FMT_PCT
                _link(celdas[4], comp.url)
                if es_exacto and precio:
                    exactos.append((precio, t))
                    relleno = ROJO if p.precio > precio + 0.5 else VERDE if p.precio < precio - 0.5 else None
                    if relleno:
                        celdas[1].fill = celdas[2].fill = relleno
                else:
                    for x in celdas[:4] + celdas[5:]:
                        x.font = GRIS_ITALICA
            else:
                nivel = c.nivel if c else "Sin datos"
                ws.cell(row=fila, column=col + 3, value=nivel).font = GRIS_ITALICA
                if c and c.nota_ia:
                    ws.cell(row=fila, column=col + 5, value=f"IA: {c.nota_ia}").font = GRIS_ITALICA
            col += len(por_tienda)
        if exactos:
            minimo, tienda_min = min(exactos)
            ws.cell(row=fila, column=col, value=minimo).number_format = FMT_Q
            ws.cell(row=fila, column=col + 1, value=tienda_min)
        pos = _posicion(p.precio, [x for x, _ in exactos])
        cp = ws.cell(row=fila, column=col + 2, value=pos)
        cp.fill = {"Más caro": ROJO, "Más barato": VERDE, "Igual al más barato": AMARILLO}.get(pos, PatternFill())

    ws.freeze_panes = "F3"
    ws.auto_filter.ref = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"
    anchos = {1: 20, 2: 55, 3: 20, 4: 15, 5: 13}
    col = len(fijas) + 1
    for _ in tiendas:
        anchos.update({col: 12, col + 1: 12, col + 2: 10, col + 3: 22, col + 4: 45, col + 5: 40})
        col += len(por_tienda)
    anchos.update({col: 14, col + 1: 15, col + 2: 18})
    _anchos(ws, anchos)
    ws.row_dimensions[2].height = 32

    # ---------------------------------------------------------------- Resumen
    wr = wb.create_sheet("Resumen")
    wr["A1"] = f"Comparación de precios — {alcance}"
    wr["A1"].font = Font(bold=True, size=14)
    wr["A2"] = f"Generado: {datetime.now():%Y-%m-%d %H:%M}   ·   Productos Imeqmo: {len(imeqmo)}"
    wr["A3"] = resumen_ia
    _encabezado(wr, 5, ["Tienda", "Productos en catálogo", "Exactos", "Similares", "Sin coincidencia",
                        "Imeqmo más barato", "Imeqmo más caro", "Diferencia promedio (exactos)", "Estado"])
    for fila, t in enumerate(tiendas, 6):
        cs = [resultados.get(p.url, {}).get(t) for p in imeqmo]
        exactos = [(p, c) for p, c in zip(imeqmo, cs) if c and c.producto and c.nivel in em.EXACTOS
                   and precio_comparable(c.producto)]
        difs = [(p.precio - precio_comparable(c.producto)) / precio_comparable(c.producto) for p, c in exactos]
        valores = [
            t, len(catalogos.get(t, [])), len(exactos),
            sum(1 for c in cs if c and c.producto and c.nivel in em.SIMILARES),
            sum(1 for c in cs if not c or not c.producto),
            sum(1 for d in difs if d < -0.0001), sum(1 for d in difs if d > 0.0001),
            (sum(difs) / len(difs)) if difs else None,
            f"No disponible: {errores[t]}" if t in errores else "OK",
        ]
        for col, v in enumerate(valores, 1):
            wr.cell(row=fila, column=col, value=v)
        wr.cell(row=fila, column=8).number_format = FMT_PCT
    fila = 6 + len(tiendas) + 1
    wr.cell(row=fila, column=1, value="Diferencia = (precio Imeqmo − precio competidor) / precio competidor. "
            "Positivo: Imeqmo más caro.").font = GRIS_ITALICA
    wr.cell(row=fila + 1, column=1, value="Solo las coincidencias 'Exacto' entran en los totales y en el mínimo "
            "de la competencia; las 'Similar' se muestran como referencia.").font = GRIS_ITALICA
    nota_intelaf = {
        "promo": "Intelaf: se compara contra su promoción general cuando existe (p. ej. 'Promo ASUS'); los "
                 "descuentos por forma de pago ('Beneficio Efectivo') se muestran en la columna Detalle.",
        "oferta": "Intelaf: se compara contra el precio más bajo que publica (incluye descuentos por forma de pago).",
        "normal": "Intelaf: se compara contra su precio normal; las ofertas se muestran en Detalle.",
    }[config.INTELAF_PRECIO]
    wr.cell(row=fila + 2, column=1, value=nota_intelaf).font = GRIS_ITALICA

    # Por categoría
    fila += 4
    wr.cell(row=fila, column=1, value="Por categoría").font = Font(bold=True, size=12)
    fila += 1
    encabezado = ["Categoría", "Productos Imeqmo"]
    for t in tiendas:
        encabezado += [f"{t}: exactos", f"{t}: similares"]
    encabezado += ["Con al menos un exacto", "Imeqmo más caro", "Imeqmo más barato"]
    _encabezado(wr, fila, encabezado)
    wr.row_dimensions[fila].height = 32
    por_categoria: dict[str, list[Producto]] = {}
    for p in imeqmo:
        por_categoria.setdefault(_categoria(p), []).append(p)
    for nombre_cat, productos in por_categoria.items():
        fila += 1
        valores = [nombre_cat, len(productos)]
        for t in tiendas:
            cs = [resultados.get(p.url, {}).get(t) for p in productos]
            valores += [sum(1 for c in cs if c and c.producto and c.nivel in em.EXACTOS),
                        sum(1 for c in cs if c and c.producto and c.nivel in em.SIMILARES)]
        posiciones = []
        for p in productos:
            mejores = [precio_comparable(c.producto) for c in resultados.get(p.url, {}).values()
                       if c.producto and c.nivel in em.EXACTOS and precio_comparable(c.producto)]
            posiciones.append(_posicion(p.precio, mejores))
        valores += [sum(1 for x in posiciones if x != "Sin comparación"),
                    posiciones.count("Más caro"), posiciones.count("Más barato")]
        for col, v in enumerate(valores, 1):
            wr.cell(row=fila, column=col, value=v)
    _anchos(wr, {1: 24, 2: 12, 3: 10, 4: 10, 5: 12, 6: 12, 7: 12, 8: 16, 9: 14, 10: 12, 11: 12})
    wr.row_dimensions[5].height = 32

    # ---------------------------------------------------------------- Detalle
    wd = wb.create_sheet("Detalle")
    _encabezado(wd, 1, ["Categoría", "Producto Imeqmo", "No. parte", "Specs Imeqmo", "Tienda", "#", "Elegido",
                        "Nivel (reglas)", "Puntaje", "Candidato", "Precio", "Specs candidato", "Motivo",
                        "Diferencias", "URL candidato"])
    fila = 2
    for p in imeqmo:
        for t in tiendas:
            c = resultados.get(p.url, {}).get(t)
            if not c:
                continue
            for i, (nivel, puntaje, comp, motivo, difs) in enumerate(c.candidatos, 1):
                valores = [_categoria(p), p.nombre, p.no_parte, _specs_txt(p.specs), t, i,
                           "Sí" if c.producto is comp else "", nivel, round(puntaje, 1), comp.nombre,
                           precio_comparable(comp), _specs_txt(comp.specs), motivo, difs, comp.url]
                for col, v in enumerate(valores, 1):
                    wd.cell(row=fila, column=col, value=v)
                wd.cell(row=fila, column=11).number_format = FMT_Q
                _link(wd.cell(row=fila, column=15), comp.url)
                fila += 1
    wd.freeze_panes = "A2"
    wd.auto_filter.ref = f"A1:O{max(fila - 1, 1)}"
    _anchos(wd, {1: 20, 2: 45, 3: 18, 4: 45, 5: 13, 6: 4, 7: 8, 8: 30, 9: 8, 10: 50, 11: 12, 12: 45, 13: 35,
                 14: 35, 15: 30})

    # ---------------------------------------------------------------- Catálogos
    wc = wb.create_sheet("Catálogos")
    _encabezado(wc, 1, ["Tienda", "Categoría", "Producto", "Precio", "Precio comparable", "Marca", "CPU",
                        "RAM (GB)", "Almac. (GB)", "Pantalla", "GPU", "UPC", "MPN / código", "Disponible", "URL"])
    fila = 2
    for t, productos in [("Imeqmo", imeqmo)] + [(t, catalogos.get(t, [])) for t in tiendas]:
        for comp in productos:
            s = comp.specs
            valores = [t, _categoria(comp), comp.nombre, comp.precio, precio_comparable(comp), s.get("marca"),
                       s.get("cpu"), s.get("ram"), s.get("almacenamiento"), s.get("pantalla"), s.get("gpu"),
                       comp.upc, comp.mpn or comp.no_parte, {True: "Sí", False: "No"}.get(comp.disponible, ""),
                       comp.url]
            for col, v in enumerate(valores, 1):
                wc.cell(row=fila, column=col, value=v)
            wc.cell(row=fila, column=4).number_format = wc.cell(row=fila, column=5).number_format = FMT_Q
            fila += 1
    wc.freeze_panes = "A2"
    wc.auto_filter.ref = f"A1:O{max(fila - 1, 1)}"
    _anchos(wc, {1: 13, 2: 20, 3: 60, 4: 12, 5: 12, 6: 12, 7: 16, 8: 8, 9: 9, 10: 8, 11: 10, 12: 15, 13: 20,
                 14: 9, 15: 40})

    wb.active = 0
    ruta.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(ruta)
    except PermissionError:  # el archivo del mismo minuto está abierto en Excel
        ruta = ruta.with_stem(f"{ruta.stem}_{datetime.now():%S}")
        wb.save(ruta)
    return ruta


def agregar_historico(ruta: Path, imeqmo: list[Producto], resultados: dict[str, dict[str, Coincidencia]]):
    nuevo = not ruta.exists()
    fecha = datetime.now().strftime("%Y-%m-%d %H:%M")
    with ruta.open("a", encoding="utf-8-sig" if nuevo else "utf-8", newline="") as f:
        w = csv.writer(f)
        if nuevo:
            w.writerow(["fecha", "categoria", "no_parte_imeqmo", "producto_imeqmo", "precio_imeqmo", "tienda",
                        "nivel", "producto_competidor", "precio_competidor", "url_competidor"])
        for p in imeqmo:
            for t, c in resultados.get(p.url, {}).items():
                comp = c.producto
                w.writerow([fecha, _categoria(p), p.no_parte, p.nombre, p.precio, t, c.nivel,
                            comp.nombre if comp else "", precio_comparable(comp) if comp else "",
                            comp.url if comp else ""])
